#!/usr/bin/env python3
"""
JBVNL electricity-bill auto-fetcher.

Logs into the JBVNL customer portal (https://cportal.jbvnl.co.in), downloads the
CURRENT month's bill PDF for each configured account, and files it into the right
property folder as `M:YY.pdf` (matching the existing naming convention).

WHY only the current month:
  The portal only serves a valid PDF for the latest bill — older months return an
  empty (0-byte) response from the server. So this is a *preventive* pipeline: run it
  monthly (right after bills issue) and you'll never miss a month again.

CREDENTIALS:
  Your password is NEVER stored in this file. It is read at runtime from the
  environment / a local .env file that YOU create (see .env.example).
  Copy .env.example -> .env and fill in JBVNL_PASSWORD.

USAGE:
  source venv/bin/activate
  python fetch_bills.py            # headless, all accounts
  HEADED=1 python fetch_bills.py   # watch it drive the browser (good first run)
  python fetch_bills.py --account 100000000000   # just one account
"""

import os, re, sys, csv, json, argparse, datetime
from pathlib import Path
from dotenv import load_dotenv
from playwright.sync_api import sync_playwright, TimeoutError as PWTimeout

sys.path.insert(0, str(Path(__file__).resolve().parent))
from notify import send as tg_send  # Telegram notification

load_dotenv()

PORTAL = "https://cportal.jbvnl.co.in/cportal/#/user/login"
HISTORY = "https://cportal.jbvnl.co.in/cportal/#/history"

# All personal config lives in .env (git-ignored) so this code is safe to publish.
BILLS_BASE = Path(os.environ.get("BILLS_BASE", str(Path.home() / "Electricity Bills")))
# Running log of every month's bill amount (even months whose PDF won't download).
# Stored as a "Bill Amounts" sheet inside the existing Electricity.xlsx (Sheet1 untouched).
XLSX_PATH = BILLS_BASE / "Electricity.xlsx"
SHEET_NAME = "Bill Amounts"
COLS = ["account", "property", "billing_date", "units_billed",
        "bill_amount", "arrears", "due_date", "scraped_at"]


def _load_accounts():
    """Parse JBVNL_ACCOUNTS='acct:Label,acct:Label' -> {acct: 'acct (Label)'}."""
    accts = {}
    for part in os.environ.get("JBVNL_ACCOUNTS", "").split(","):
        num, _, label = part.strip().partition(":")
        num, label = num.strip(), label.strip()
        if num:
            accts[num] = f"{num} ({label})" if label else num
    return accts


# account number -> destination folder name (inside BILLS_BASE)
ACCOUNTS = _load_accounts()

# Password source (per-account override supported: JBVNL_PW_<account>), else shared.
def password_for(acct: str) -> str:
    return os.environ.get(f"JBVNL_PW_{acct}") or os.environ.get("JBVNL_PASSWORD") or ""


def dest_name_from_download(filename: str) -> str | None:
    """`ACCOUNT_62026_E.pdf` -> `6:26.pdf`  (month 6, year 2026)."""
    m = re.search(r"_(\d{1,2})(\d{4})_", filename)
    if not m:
        return None
    month = int(m.group(1))
    yy = m.group(2)[-2:]
    return f"{month}:{yy}.pdf"


def _num(s: str):
    """'₹40,364.16' -> 40364.16 ; '' -> ''"""
    s = re.sub(r"[^\d.\-]", "", s or "")
    try:
        return float(s) if s else ""
    except ValueError:
        return ""


def scrape_bill_table(table, acct: str, folder: str):
    """Read every month's row (billing date, units, amount, arrears, due date)."""
    rows = []
    trs = table.locator("tr", has=table.page.locator("td"))
    for i in range(trs.count()):
        cells = trs.nth(i).locator("td")
        vals = [cells.nth(j).inner_text().strip() for j in range(cells.count())]
        # expected order: Billing Date, Units, Bill Amount, Arrears, Due Date, [Detailed Bill]
        if len(vals) < 5:
            continue
        rows.append({
            "account": acct,
            "property": re.sub(r"^\d+\s*\(|\)$", "", folder).strip(),
            "billing_date": vals[0],
            "units_billed": _num(vals[1]),
            "bill_amount": _num(vals[2]),
            "arrears": _num(vals[3]),
            "due_date": vals[4],
            "scraped_at": datetime.date.today().isoformat(),
        })
    return rows


def _sortkey(r):
    try:
        d = datetime.datetime.strptime(str(r["billing_date"]), "%d-%b-%Y")
    except Exception:
        d = datetime.datetime.min
    return (str(r.get("account", "")), -d.timestamp())


def store_amounts(new_rows):
    """Merge scraped rows into the 'Bill Amounts' sheet of Electricity.xlsx,
    keyed by (account, billing_date). Sheet1 (and any other sheet) is untouched."""
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment

    if XLSX_PATH.exists():
        wb = openpyxl.load_workbook(XLSX_PATH)
    else:
        wb = openpyxl.Workbook()
        wb.remove(wb.active)

    existing = {}
    if SHEET_NAME in wb.sheetnames:
        ws_old = wb[SHEET_NAME]
        headers = [c.value for c in ws_old[1]]
        for row in ws_old.iter_rows(min_row=2, values_only=True):
            rec = dict(zip(headers, row))
            if rec.get("billing_date"):
                existing[(str(rec.get("account")), rec.get("billing_date"))] = rec
        idx = wb.sheetnames.index(SHEET_NAME)
        wb.remove(ws_old)
    else:
        idx = len(wb.sheetnames)  # append at end

    for r in new_rows:
        existing[(str(r["account"]), r["billing_date"])] = r

    ws = wb.create_sheet(SHEET_NAME, idx)
    ws.append(COLS)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="305496")
        c.alignment = Alignment(horizontal="center")
    ws.freeze_panes = "A2"
    for r in sorted(existing.values(), key=_sortkey):
        ws.append([r.get(k, "") for k in COLS])
    # currency-ish number format + column widths
    for col in ("E", "F"):  # bill_amount, arrears
        for cell in ws[col][1:]:
            cell.number_format = "#,##0.00"
    widths = {"A": 13, "B": 10, "C": 14, "D": 12, "E": 13, "F": 13, "G": 14, "H": 12}
    for c, w in widths.items():
        ws.column_dimensions[c].width = w

    wb.save(XLSX_PATH)
    return len(existing)


def fetch_account(page, acct: str, folder: str, pw: str) -> str:
    dest_dir = BILLS_BASE / folder
    dest_dir.mkdir(parents=True, exist_ok=True)

    # --- login ---
    page.goto(PORTAL, wait_until="networkidle")
    page.fill('input[formcontrolname="customerNo"]', acct)
    page.fill('input[formcontrolname="password"]', pw)
    page.click('button:has-text("LETS GO")')

    # wait until we leave the login route
    try:
        page.wait_for_url("**/#/home", timeout=20000)
    except PWTimeout:
        # some portals land elsewhere; verify we're authenticated by absence of login form
        if page.query_selector('input[formcontrolname="password"]'):
            raise RuntimeError(f"[{acct}] login failed — check password / account number")

    # --- go to bill history, table view ---
    page.goto(HISTORY, wait_until="networkidle")
    page.wait_for_timeout(1500)
    # click the 'table_chart' toggle on the Bill History card (first occurrence)
    page.get_by_text("table_chart", exact=True).first.click()
    page.wait_for_timeout(1200)

    # --- find the Bill History table and its first data row ---
    table = page.locator("table", has_text="Detailed Bill").first
    table.wait_for(timeout=10000)

    # --- scrape ALL months' amounts (works even for months whose PDF won't download) ---
    scraped = scrape_bill_table(table, acct, folder)
    if scraped:
        total = store_amounts(scraped)
        print(f"[{acct}] logged {len(scraped)} months to Electricity.xlsx → '{SHEET_NAME}' ({total} rows total)")

    # first DATA row = current month (the only month the portal serves a real PDF for).
    # We filter to rows that contain <td> (header uses <th>), so index 0 is the top data row.
    row = table.locator("tr", has=page.locator("td")).nth(0)
    row.wait_for(timeout=10000)
    last_cell = row.locator("td").last
    pdf_icon = last_cell.locator("span").first  # first span = the "_E" bill

    # --- download ---
    with page.expect_download(timeout=30000) as dl_info:
        pdf_icon.click()
    download = dl_info.value
    suggested = download.suggested_filename

    dest_name = dest_name_from_download(suggested) or f"latest-{datetime.date.today()}.pdf"
    dest_path = dest_dir / dest_name
    download.save_as(str(dest_path))

    size = dest_path.stat().st_size
    if size < 1000:
        dest_path.unlink(missing_ok=True)
        raise RuntimeError(
            f"[{acct}] portal returned an empty PDF ({size} B) — the current bill "
            f"may not be generated yet. Try again after the bill is issued."
        )
    return f"[{acct}] saved {folder}/{dest_name}  ({size:,} bytes)"


def logout(page):
    # clearing storage + going to login drops the session between accounts
    try:
        page.context.clear_cookies()
        page.evaluate("() => { sessionStorage.clear(); localStorage.clear(); }")
    except Exception:
        pass


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--account", help="only fetch this account number")
    args = ap.parse_args()

    targets = {args.account: ACCOUNTS[args.account]} if args.account else ACCOUNTS
    headed = os.environ.get("HEADED") == "1"

    results, errors = [], []
    with sync_playwright() as p:
        browser = p.chromium.launch(headless=not headed)
        for acct, folder in targets.items():
            pw = password_for(acct)
            if not pw:
                errors.append(f"[{acct}] no password — set JBVNL_PASSWORD in .env")
                continue
            context = browser.new_context(accept_downloads=True)
            page = context.new_page()
            try:
                results.append(fetch_account(page, acct, folder, pw))
            except Exception as e:
                errors.append(str(e))
                try:
                    page.screenshot(path=f"error_{acct}.png")
                    errors.append(f"    (screenshot: error_{acct}.png)")
                except Exception:
                    pass
            finally:
                logout(page)
                context.close()
        browser.close()

    print("\n=== JBVNL bill fetch ===")
    for r in results:
        print(" ✅", r)
    for e in errors:
        print(" ⚠️ ", e)

    # --- Telegram summary ---
    stamp = datetime.datetime.now().strftime("%d %b %Y, %H:%M")
    lines = [f"🔌 <b>JBVNL bill fetch</b> — {stamp}"]
    for r in results:
        lines.append(f"✅ {r}")
    for e in errors:
        lines.append(f"⚠️ {e}")
    if not results and not errors:
        lines.append("No accounts processed.")
    tg_send("\n".join(lines))

    sys.exit(1 if errors and not results else 0)


if __name__ == "__main__":
    main()
